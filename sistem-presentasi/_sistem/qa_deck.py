# -*- coding: utf-8 -*-
"""QA otomatis per-produksi (04). Pakai: python3 qa_deck.py <pptx> [min_slides]
Exit 0 = lulus cek otomatis; 1 = ada kegagalan. (Cek manusia tetap di 04.B.)"""
import sys
from pptx import Presentation
from pptx.oxml.ns import qn

def hastash(t): return any(0x064B<=ord(c)<=0x065F or ord(c)==0x0670 or 0x06D6<=ord(c)<=0x06ED for c in t)
BAD="()—+"

def main():
    path=sys.argv[1]
    min_sl=int(sys.argv[2]) if len(sys.argv)>2 else 1
    prs=Presentation(path)
    fails=[]; warns=[]
    n=len(prs.slides._sldIdLst)
    print("slide:",n)
    if n<min_sl: fails.append(f"jumlah slide {n} < rencana {min_sl}")
    for i,sl in enumerate(prs.slides,1):
        rtl=False; ea=False; notes=False
        for sh in sl.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    pPr=p._p.find(qn('a:pPr'))
                    if pPr is not None and pPr.get('rtl')=='1': rtl=True
                    for r in p.runs:
                        rPr=r._r.find(qn('a:rPr'))
                        if rPr is not None and rPr.find(qn('a:ea')) is not None: ea=True
            if sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        if hastash(c.text) or any(b in c.text for b in BAD): fails.append(f"S{i} tabel: karakter bermasalah")
        try: notes=bool(sl.notes_slide.notes_text_frame.text.strip())
        except Exception: notes=False
        if not rtl: warns.append(f"S{i} tanpa rtl")
        if not ea: warns.append(f"S{i} tanpa ea-font")
        if not notes: fails.append(f"S{i} tanpa notes")
    # gambar & crop
    pics=0
    for i,sl in enumerate(prs.slides,1):
        for sh in sl.shapes:
            if sh.shape_type==13:
                pics+=1
                cr=getattr(sh,'crop_right',0) or 0
                if cr==0: warns.append(f"S{i} gambar tanpa crop (cek gepeng bila kotak non-proporsional)")
    print("gambar:",pics)
    # tashkeel global
    for i,sl in enumerate(prs.slides,1):
        for sh in sl.shapes:
            t = sh.text_frame.text if sh.has_text_frame else ""
            if hastash(t): fails.append(f"S{i} masih ada tashkeel")
            if any(b in t for b in BAD): fails.append(f"S{i} ada karakter bidi-riskan")
    print("FAIL:",fails if fails else "TIDAK ADA")
    print("WARN:",warns if warns else "TIDAK ADA")
    sys.exit(1 if fails else 0)

if __name__=="__main__":
    main()
