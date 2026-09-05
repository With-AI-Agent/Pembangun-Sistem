# -*- coding: utf-8 -*-
"""v2 — deck tesis fikih hiasan wanita.
Perbaikan dari v1: (1) desain lebih profesional: font besar, isi terpusat vertikal,
footer bar hijau + nomor slide Arab, spine kanan hijau-emas; (2) isi diimbang ke
HASIL/temuan di depan (dari ملخص hal 6), metode dipadatkan.
Semua konten berjejak halaman; tidak ada karangan.
"""
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

GREEN=RGBColor(0x0F,0x3D,0x2E); GOLD=RGBColor(0xC9,0xA2,0x27)
CREAM=RGBColor(0xF7,0xF3,0xE9); INK=RGBColor(0x22,0x22,0x22)
DARKGREEN=RGBColor(0x0B,0x2E,0x23); FONT="Amiri"
W=13.333; H=7.5
AD="٠١٢٣٤٥٦٧٨٩"
def anum(n): return "".join(AD[int(d)] for d in str(n))

prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H)
BLANK=prs.slide_layouts[6]

def rrun(run,size,bold,color):
    run.font.name=FONT; run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=color
    rPr=run._r.get_or_add_rPr()
    for tag in ("a:ea","a:cs"):
        el=rPr.find(qn(tag))
        if el is None: el=rPr.makeelement(qn(tag),{}); rPr.append(el)
        el.set("typeface",FONT)

def rpara(p,size,bold,color,after=10,space=1.15):
    p.alignment=PP_ALIGN.RIGHT
    pPr=p._p.get_or_add_pPr(); pPr.set("rtl","1")
    p.space_after=Pt(after)
    try: p.line_spacing=space
    except Exception: pass
    if p.runs: rrun(p.runs[0],size,bold,color)

def bg(s):
    f=s.background.fill; f.solid(); f.fore_color.rgb=CREAM

def box(s,l,t,w,h): return s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))

def new_slide(num, src):
    s=prs.slides.add_slide(BLANK); bg(s)
    # spine kanan: hijau + garis emas
    sp=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(W-0.30),0,Inches(0.30),Inches(H))
    sp.fill.solid(); sp.fill.fore_color.rgb=GREEN; sp.line.fill.background()
    gl=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(W-0.36),0,Inches(0.06),Inches(H))
    gl.fill.solid(); gl.fill.fore_color.rgb=GOLD; gl.line.fill.background()
    # footer bar
    fb=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,Inches(H-0.55),Inches(W),Inches(0.55))
    fb.fill.solid(); fb.fill.fore_color.rgb=GREEN; fb.line.fill.background()
    # nomor slide kiri
    nb=box(s,0.5,H-0.5,2.0,0.45); p=nb.text_frame.paragraphs[0]; p.text=anum(num)
    p.alignment=PP_ALIGN.LEFT; rrun(p.runs[0],16,True,GOLD)
    # sumber kanan
    sb=box(s,W-3.0,H-0.5,2.4,0.45); p=sb.text_frame.paragraphs[0]; p.text=src
    p.alignment=PP_ALIGN.RIGHT; rrun(p.runs[0],14,False,GOLD)
    return s

def title(s,text,size=34):
    tb=box(s,0.7,0.45,W-1.6,1.3); tf=tb.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.text=text; rpara(p,size,True,GREEN,after=6)
    line=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(W-0.36-6.4),Inches(1.75),Inches(6.4),Inches(0.06))
    line.fill.solid(); line.fill.fore_color.rgb=GOLD; line.line.fill.background()

def bullets(s,items,size=26):
    tb=box(s,0.9,2.1,W-2.0,4.5); tf=tb.text_frame; tf.word_wrap=True
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    for i,item in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=item; rpara(p,size,False,INK,after=18,space=1.2)

def notes(s,t): s.notes_slide.notes_text_frame.text=t

def table_slide(s,rows):
    sh=s.shapes.add_table(len(rows),2,Inches(1.1),Inches(2.1),Inches(W-2.6),Inches(4.3))
    t=sh.table; t.columns[0].width=Inches(3.2); t.columns[1].width=Inches(W-2.6-3.2)
    for r,(a,b) in enumerate(rows):
        for c,val in ((0,a),(1,b)):
            cell=t.cell(r,c); cell.text=val
            p=cell.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.RIGHT
            p._p.get_or_add_pPr().set("rtl","1")
            run=p.runs[0]; run.font.name=FONT; run.font.size=Pt(20 if r else 22); run.font.bold=(r==0)
            rPr=run._r.get_or_add_rPr()
            for tag in ("a:ea","a:cs"):
                el=rPr.find(qn(tag)) or rPr.makeelement(qn(tag),{})
                if el not in rPr: rPr.append(el)
                el.set("typeface",FONT)
            if r==0: cell.fill.solid(); cell.fill.fore_color.rgb=GREEN; run.font.color.rgb=CREAM
            elif r%2==0: cell.fill.solid(); cell.fill.fore_color.rgb=RGBColor(0xEC,0xE6,0xD6)
            else: cell.fill.solid(); cell.fill.fore_color.rgb=CREAM

# ---- S1 judul ----
s=new_slide(1,"ص٦")
tb=box(s,1.2,2.0,W-2.4,2.6); tf=tb.text_frame; tf.word_wrap=True
p=tf.paragraphs[0]; p.text="أحكام النوازل الفقهية المعاصرة المتعلقة بزينة المرأة في المذاهب الأربعة"
rpara(p,40,True,GREEN,after=20,space=1.2)
p2=tf.add_paragraph(); p2.text="إعداد: الشيخ علي بلعيدي"; rpara(p2,26,False,INK)
# bingkai emas
fr=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(1.0),Inches(1.7),Inches(W-2.0),Inches(3.4))
fr.fill.background(); fr.line.color.rgb=GOLD; fr.line.width=Pt(2)
notes(s,"Slide pembuka. Nama pembimbing/universitas tidak ada di bahan -> kosong, jangan dikarang.")

# ---- S2 temuan utama ----
s=new_slide(2,"ص٦"); title(s,"تدور أحكام الزينة المعاصرة على ثلاث علل: تغيير خلق الله، التدليس، والضرر")
bullets(s,[
 "الأصل: التحريم عند تحقق إحدى العلل الثلاث",
 "الجواز مشروط بانتفاء التغيير والتدليس والضرر",
 "المنهج: وصفي مقارن، استقرائي، استنباطي",
])
notes(s,"Temuan utama (hal 6): hukum berputar pada tiga illat.")

# ---- S3 ضوابط per jenis ----
s=new_slide(3,"ص٦"); title(s,"النمص والوشم والتشقير ونحوها تمتنع عند قيام التغيير أو التدليس")
bullets(s,[
 "الوصل: محرم، ويجوز بغير شعر الآدمي اتفاقًا",
 "الضابط: كلما تحققت علة محرِّمة امتنع الفعل",
])
notes(s,"Hal 6: jenis hiasan & ضابط-nya; وصل boleh dengan bukan rambut manusia.")

# ---- S4 operasi + efek nikah ----
s=new_slide(4,"ص٦"); title(s,"العمليات تجوز بضوابط، وللزينة أثر في النظر داخل عقد النكاح")
bullets(s,[
 "العمليات الجراحية تجوز للحاجة الشرعية دون تغيير الخلق",
 "الزينة المؤثرة على النظر تُرتَّب أحكامها في النكاح وفق المذاهب",
])
notes(s,"Hal 6: operasi boleh bila kebutuhan sah; hiasan memengaruhi pandangan dalam akad nikah.")

# ---- S5 tujuan ----
s=new_slide(5,"ص٢١"); title(s,"أربعة أهداف: الأدوات، العمليات، البيع، وأثر الزينة على النكاح")
bullets(s,[
 "بيان حكم أدوات التجميل المعاصرة في المذاهب الأربعة",
 "معرفة حكم عمليات التجميل في المذاهب المعتمدة",
 "معرفة حكم بيع أدوات التجميل في المذاهب الأربعة",
 "بيان أثر الزينة على عقد النكاح",
],size=24)
notes(s,"Tujuan (hal 21).")

# ---- S6 manfaat ----
s=new_slide(6,"ص٢١"); title(s,"الفوائد توازي الأهداف وتُعمّق إدراك الأحكام")
bullets(s,[
 "معرفة أحكام الزينة المعاصرة في المذاهب الأربعة",
 "بيان أحكام النوازل المتعلقة بزينة المرأة",
 "معرفة حكم العمليات وبيع الأدوات الممنوعة",
 "إدراك أثر الزينة على عقد النكاح",
],size=24)
notes(s,"Manfaat (hal 21).")

# ---- S7 masalah ----
s=new_slide(7,"ص٢١"); title(s,"خمس مشكلات استدعت البحث: نوازل، حصر، أزمنة، واردات، تكييف")
bullets(s,[
 "نوازل معاصرة لم يفصّلها الفقهاء قديمًا",
 "صعوبة حصر وسائل الزينة لاختلاط العادات",
 "اختلاف الأزمنة يستدعي نظر القواعد والضوابط",
 "ورود أدوات من الغرب دون نظر شرعي واختلاف التكييف",
],size=24)
notes(s,"Masalah (hal 21).")

# ---- S8 metode & jenis ----
s=new_slide(8,"ص٢٣·٨٨"); title(s,"بحث فقهي مقارن بثلاثة مناهج: وصفي، استقرائي، استنباطي")
bullets(s,[
 "وصفي مقارن: تتبّع أقوال المذاهب ومقارنتها",
 "استقرائي: استقراء النوازل والدراسات المعاصرة",
 "استنباطي: استنباط الحكم من النصوص والقواعد",
])
notes(s,"Metode (hal 23) + jenis (hal 88) dipadatkan.")

# ---- S9 prosedur data ----
s=new_slide(9,"ص٨٩"); title(s,"جمعٌ ومعالجةٌ وتحقق: من التتبّع إلى التطبيق الموثّق")
bullets(s,[
 "جمع: تتبّع أقوال المذاهب وجمع النوازل وترتيبها",
 "معالجة: مقارنة الأقوال وذكر الأدلة ثم التطبيق",
 "تحقق: عرض المعلومة على عدة مصادر حتى يتفق المعنى",
])
notes(s,"Prosedur data (hal 89) dipadatkan: Far'2+3+4.")

# ---- S10 علة table ----
s=new_slide(10,"ص٦٦–٦٧"); title(s,"اتفقوا على الجواز بغير شعر الآدمي، واختلفوا في العلة — والراجح التدليس")
table_slide(s,[
 ("المذهب","العلة في تحريم وصل الشعر"),
 ("الحنفية","كرامة الإنسان وتحريم الابتذال"),
 ("المالكية","مركّبة: تغيير خلق الله + التدليس"),
 ("الشافعية","مركّبة: الكرامة + تحريم استعمال النجس"),
 ("الحنابلة","التدليس والغش"),
 ("الراجح","الزور والتدليس (حديث المرأة)"),
])
notes(s,"العلة (hal 66-67). Tabel perbandingan mazhab.")

# ---- S11 kesimpulan ----
s=new_slide(11,"ص٦"); title(s,"الخلاصة: التحريم عند العلل الثلاث، والجواز بضوابط، وللزينة أثر في النكاح")
bullets(s,[
 "التحريم عند تغيير الخلق أو التدليس أو الضرر",
 "جواز العمليات والوصل بغير الآدمي بضوابطها",
 "ترتيب أثر الزينة على النظر في عقد النكاح",
])
notes(s,"Kesimpulan dirangkum dari ملخص hal 6; خاتمة hal 151 = doa.")

# ---- S12 توصيات ----
s=new_slide(12,"ص١٥٢"); title(s,"ثلاث وصايا: للعلماء، للمرأة المسلمة، وللمجامع الفقهية")
bullets(s,[
 "للعلماء: دقة البحث وإظهار أحكام النوازل",
 "للمرأة المسلمة: الحذر من المحرم والاكتفاء بالحلال عن علم",
 "للمجامع الفقهية: متابعة المستجدات وبيان أحكامها",
])
notes(s,"توصيات (hal 152).")

# ---- S13 penutup ----
s=new_slide(13,"ص١٥١"); title(s,"الحمد لله؛ جهد بشري يقبل الخطأ، وسُئل الله القبول والنفع")
bullets(s,[
 "الحمد لله على إتمام البحث",
 "سؤال الله القبول والنفع به",
])
notes(s,"خاتمة (hal 151) = doa penutup.")

OUT="sistem-presentasi/deck-aktif/presentasi-tesis-fikih-hiasan-wanita/keluaran/presentasi-tesis-fikih-hiasan-wanita.pptx"
prs.save(OUT)
print("saved v2 slides:", len(prs.slides._sldIdLst))
