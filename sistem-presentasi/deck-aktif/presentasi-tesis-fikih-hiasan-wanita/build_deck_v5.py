# -*- coding: utf-8 -*-
"""v5 — sesuai _sistem/02: (G) sanitizer tashkeel otomatis + frasa sederhana;
(H) layout kaya: slide metode jadi DASHBOARD 3 kartu ROUNDED (multi-objek).
Aspect-safe band & bebas-kurung dipertahankan dari v4.
"""
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os, struct

GREEN=RGBColor(0x0F,0x3D,0x2E); GOLD=RGBColor(0xC9,0xA2,0x27)
CREAM=RGBColor(0xF7,0xF3,0xE9); INK=RGBColor(0x22,0x22,0x22)
WHITE=RGBColor(0xFF,0xFF,0xFF)
FONT="Amiri"; W=13.333; H=7.5
HERE=os.path.dirname(os.path.abspath(__file__))
IMG=os.path.join(HERE,"gambar","pola-geometris.png")
_IW,_IH=(1376,768)
if os.path.exists(IMG):
    _d=open(IMG,'rb').read(); _IW,_IH=struct.unpack(">II",_d[16:24])
AD="٠١٢٣٤٥٦٧٨٩"
def anum(n): return "".join(AD[int(d)] for d in str(n))
def clean(t):
    return "".join(ch for ch in t
        if not (0x064B<=ord(ch)<=0x065F) and ord(ch)!=0x0670 and not (0x06D6<=ord(ch)<=0x06ED))

prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H)
BLANK=prs.slide_layouts[6]

def rrun(run,size,bold,color):
    run.font.name=FONT; run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=color
    rPr=run._r.get_or_add_rPr()
    for tag in ("a:ea","a:cs"):
        el=rPr.find(qn(tag))
        if el is None: el=rPr.makeelement(qn(tag),{}); rPr.append(el)
        el.set("typeface",FONT)

def rpara(p,size,bold,color,after=10,space=1.15,level=0):
    p.alignment=PP_ALIGN.RIGHT
    pPr=p._p.get_or_add_pPr(); pPr.set("rtl","1")
    p.level=level; p.space_after=Pt(after)
    try: p.line_spacing=space
    except Exception: pass
    if p.runs: rrun(p.runs[0],size,bold,color)

def bg(s):
    f=s.background.fill; f.solid(); f.fore_color.rgb=CREAM
def box(s,l,t,w,h): return s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))

def add_band(s,band_w):
    if not os.path.exists(IMG): return
    f=(band_w/H)*_IH/_IW
    if f>1: f=1
    pic=s.shapes.add_picture(IMG,Inches(W-band_w),0,Inches(band_w),Inches(H))
    pic.crop_left=0.0; pic.crop_top=0.0; pic.crop_bottom=0.0; pic.crop_right=1-f

def new_slide(num,src,band_w=0.9):
    s=prs.slides.add_slide(BLANK); bg(s)
    add_band(s,band_w)
    gl=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(W-band_w-0.06),0,Inches(0.06),Inches(H))
    gl.fill.solid(); gl.fill.fore_color.rgb=GOLD; gl.line.fill.background()
    fb=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,Inches(H-0.55),Inches(W),Inches(0.55))
    fb.fill.solid(); fb.fill.fore_color.rgb=GREEN; fb.line.fill.background()
    nb=box(s,0.5,H-0.5,2.0,0.45); p=nb.text_frame.paragraphs[0]; p.text=anum(num)
    p.alignment=PP_ALIGN.LEFT; rrun(p.runs[0],16,True,GOLD)
    sb=box(s,W-3.4,H-0.5,2.6,0.45); p=sb.text_frame.paragraphs[0]; p.text=src
    p.alignment=PP_ALIGN.RIGHT; rrun(p.runs[0],14,False,GOLD)
    return s

def title(s,text,size=32):
    tb=box(s,0.7,0.45,W-2.2,1.4); tf=tb.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.text=clean(text); rpara(p,size,True,GREEN,after=6)
    line=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(W-0.96-6.4),Inches(1.8),Inches(6.4),Inches(0.06))
    line.fill.solid(); line.fill.fore_color.rgb=GOLD; line.line.fill.background()

def para(s,text,size=24):
    tb=box(s,0.9,2.1,W-2.3,4.4); tf=tb.text_frame; tf.word_wrap=True
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.text=clean(text)
    rpara(p,size,False,INK,after=0,space=1.4)

def bullets(s,items,size=24):
    tb=box(s,0.9,2.0,W-2.3,4.6); tf=tb.text_frame; tf.word_wrap=True
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    first=True
    for item in items:
        text,level = item if isinstance(item,tuple) else (item,0)
        p=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
        p.text=clean(text)
        rpara(p,size if level==0 else size-4,False,INK if level==0 else RGBColor(0x44,0x44,0x44),
              after=14 if level==0 else 8,space=1.15,level=level)

def dashboard(s,cards):
    n=len(cards); gap=0.3
    avail=W-2.3; cw=(avail-gap*(n-1))/n
    xr=W-1.4
    for i,(head,desc) in enumerate(cards):
        x=xr-cw-i*(cw+gap)
        sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(2.4),Inches(cw),Inches(3.4))
        try: sh.adjustments[0]=0.06
        except Exception: pass
        sh.fill.solid(); sh.fill.fore_color.rgb=WHITE
        sh.line.color.rgb=GOLD; sh.line.width=Pt(1.5)
        tf=sh.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
        p=tf.paragraphs[0]; p.text=clean(head); rpara(p,24,True,GREEN,after=10)
        p2=tf.add_paragraph(); p2.text=clean(desc); rpara(p2,16,False,INK,after=0,space=1.2)
        # inset margin
        tf.margin_left=Inches(0.2); tf.margin_right=Inches(0.2)

def notes(s,t): s.notes_slide.notes_text_frame.text=t

def table_slide(s,rows):
    sh=s.shapes.add_table(len(rows),2,Inches(1.1),Inches(2.1),Inches(W-3.4),Inches(4.3))
    t=sh.table; t.columns[0].width=Inches(3.0); t.columns[1].width=Inches(W-3.4-3.0)
    for r,(a,b) in enumerate(rows):
        for c,val in ((0,a),(1,b)):
            cell=t.cell(r,c); cell.text=clean(val)
            p=cell.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.RIGHT
            p._p.get_or_add_pPr().set("rtl","1")
            run=p.runs[0]; run.font.name=FONT; run.font.size=Pt(18 if r else 20); run.font.bold=(r==0)
            rPr=run._r.get_or_add_rPr()
            for tag in ("a:ea","a:cs"):
                el=rPr.find(qn(tag)) or rPr.makeelement(qn(tag),{})
                if el not in rPr: rPr.append(el)
                el.set("typeface",FONT)
            if r==0: cell.fill.solid(); cell.fill.fore_color.rgb=GREEN; run.font.color.rgb=CREAM
            elif r%2==0: cell.fill.solid(); cell.fill.fore_color.rgb=RGBColor(0xEC,0xE6,0xD6)
            else: cell.fill.solid(); cell.fill.fore_color.rgb=CREAM

# S1
s=new_slide(1,"ص٦",band_w=1.4)
tb=box(s,1.2,2.0,W-3.2,2.6); tf=tb.text_frame; tf.word_wrap=True
p=tf.paragraphs[0]; p.text=clean("أحكام النوازل الفقهية المعاصرة المتعلقة بزينة المرأة في المذاهب الأربعة")
rpara(p,38,True,GREEN,after=20,space=1.2)
p2=tf.add_paragraph(); p2.text=clean("إعداد: الشيخ علي بلعيدي"); rpara(p2,26,False,INK)
fr=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(1.0),Inches(1.7),Inches(W-3.0),Inches(3.4))
try: fr.adjustments[0]=0.04
except Exception: pass
fr.fill.background(); fr.line.color.rgb=GOLD; fr.line.width=Pt(2)
notes(s,"Pembuka. Pembimbing/universitas kosong - jangan dikarang.")

# S2 paragraf
s=new_slide(2,"ص٦"); title(s,"تدور أحكام الزينة المعاصرة حول ثلاث علل: تغيير خلق الله، والتدليس، والضرر")
para(s,"تُقاس وسائل الزينة المعاصرة بثلاث علل: تغيير خلق الله، والتدليس، والضرر؛ فمتى تحققت إحداها امتنع الفعل، ومتى انتفت جاز. ومن هنا يُعلم حكم النمص والوشم والتشقير والوصل والعدسات، والمنهج وصفي مقارن استقرائي استنباطي.")
notes(s,"Temuan utama hal 6 sebagai paragraf; harakat di-strip sanitizer.")

# S3
s=new_slide(3,"ص٦"); title(s,"لكل وسيلة زينة ضابطها: الامتناع عند قيام علة محرمة")
bullets(s,[
 "التشقير والحواجب والنمص: تمتنع عند التغيير",
 "الوشم: ممنوع لتغيير الخلقة",
 "العدسات والأظافر: تقاس بضوابط التغيير والتدليس والضرر",
 "الوصل: محرم، ويجوز بغير شعر الآدمي اتفاقا",
])
notes(s,"Hal 6.")

# S4
s=new_slide(4,"ص٦"); title(s,"العمليات الجراحية التجميلية تجوز للحاجة الشرعية دون تغيير الخلق")
bullets(s,[
 "الجواز مشروط بالحاجة الشرعية كالعلاج وإصلاح العيب",
 "يمتنع ما كان لمجرد تغيير خلق الله",
 "يراعى انتفاء الضرر",
])
notes(s,"Hal 6.")

# S5
s=new_slide(5,"ص٦"); title(s,"للزينة أثر في النظر داخل عقد النكاح، وتترتب عليه أحكام وفق المذاهب")
bullets(s,[
 "تترتب على الزينة المؤثرة على النظر أحكام في باب النكاح",
 "تفصيل المذاهب في أثرها على النظر والخيار",
])
notes(s,"Hal 6.")

# S6
s=new_slide(6,"ص٢١"); title(s,"أربعة أهداف: الأدوات، العمليات، البيع، وأثر الزينة على النكاح")
bullets(s,[
 "بيان حكم أدوات التجميل المعاصرة للمرأة في المذاهب الأربعة",
 "معرفة حكم عمليات التجميل في المذاهب المعتمدة",
 "معرفة حكم بيع أدوات التجميل في المذاهب الأربعة",
 "بيان الأثر المترتب على الزينة في عقد النكاح",
],size=22)
notes(s,"Tujuan hal 21.")

# S7
s=new_slide(7,"ص٢١"); title(s,"الفوائد توازي الأهداف وتعمق إدراك الأحكام والنوازل")
bullets(s,[
 "معرفة أحكام الزينة المعاصرة في المذاهب الأربعة",
 "بيان أحكام النوازل المتعلقة بزينة المرأة",
 "معرفة حكم عمليات التجميل وبيع الأدوات الممنوعة",
 "إدراك أثر الزينة على عقد النكاح",
],size=22)
notes(s,"Manfaat hal 21.")

# S8
s=new_slide(8,"ص٢١"); title(s,"خمس مشكلات استدعت البحث: نوازل، حصر، أزمنة، واردات، تكييف")
bullets(s,[
 "نوازل معاصرة لم يفصلها الفقهاء قديما",
 "صعوبة حصر وسائل الزينة لاختلاط العادات",
 "اختلاف الأزمنة يستدعي نظر القواعد والضوابط",
 "ورود أدوات من الغرب دون نظر شرعي",
 "اختلاف التكييف الفقهي لوسائل الزينة",
],size=22)
notes(s,"Masalah hal 21.")

# S9 DASHBOARD metode
s=new_slide(9,"ص٢٣·٨٨"); title(s,"بحث فقهي مقارن بثلاثة مناهج متكاملة")
dashboard(s,[
 ("وصفي مقارن","تتبع أقوال المذاهب من كتبها المعتمدة ومقارنتها"),
 ("استقرائي","استقراء النوازل والدراسات المعاصرة"),
 ("استنباطي","استنباط الحكم من النصوص والقواعد والضوابط"),
])
notes(s,"Metode hal 23 sebagai dashboard 3 kartu rounded (multi-objek).")

# S10
s=new_slide(10,"ص٨٩"); title(s,"جمع ومعالجة وتحقق: من التتبع إلى التطبيق الموثق")
bullets(s,[
 "جمع: تتبع أقوال المذاهب وجمع النوازل ثم الترتيب والمقارنة",
 "معالجة: مقارنة الأقوال وذكر الأدلة ثم التطبيق على النوازل",
 "تحقق: عرض المعلومة على عدة مصادر حتى يتفق المعنى",
])
notes(s,"Prosedur hal 89.")

# S11 table
s=new_slide(11,"ص٦٦–٦٧"); title(s,"اتفقوا على الجواز بغير شعر الآدمي واختلفوا في العلة، والراجح التدليس")
table_slide(s,[
 ("المذهب","العلة في تحريم وصل الشعر"),
 ("الحنفية","كرامة الإنسان وتحريم الابتذال"),
 ("المالكية","مركبة: تغيير خلق الله والتدليس"),
 ("الشافعية","مركبة: الكرامة وتحريم استعمال النجس"),
 ("الحنابلة","التدليس والغش"),
 ("الراجح","الزور والتدليس؛ لحديث المرأة"),
])
notes(s,"العلة hal 66-67.")

# S12 paragraf
s=new_slide(12,"ص٦"); title(s,"الخلاصة: التحريم عند العلل الثلاث، والجواز بضوابط، وللزينة أثر في النكاح")
para(s,"خلاصة البحث أن التحريم يدور مع علله الثلاث وجودا وعدما، وأن العمليات تجوز للحاجة والوصل يجوز بغير شعر الآدمي بضوابطهما، وأن للزينة أثرا في النظر تترتب عليه أحكام في عقد النكاح وفق المذاهب.")
notes(s,"Kesimpulan paragraf.")

# S13
s=new_slide(13,"ص١٥٢"); title(s,"ثلاث وصايا: للعلماء، للمرأة المسلمة، وللمجامع الفقهية")
bullets(s,[
 "للعلماء وطلاب العلم: دقة البحث وإظهار أحكام النوازل للمجتمع",
 "للمرأة المسلمة: الحذر من الزينة المحرمة والاكتفاء بالحلال عن علم",
 "للمجامع الفقهية: متابعة المستجدات في زينة المرأة وبيان أحكامها",
])
notes(s,"توصيات hal 152.")

# S14
s=new_slide(14,"ص١٥١"); title(s,"الحمد لله؛ جهد بشري يقبل الخطأ، ونسأل الله القبول والنفع")
bullets(s,[
 "الحمد لله على إتمام البحث",
 "نسأل الله القبول والنفع به",
])
notes(s,"خاتمة hal 151.")

OUT=os.path.join(HERE,"keluaran","presentasi-tesis-fikih-hiasan-wanita.pptx")
prs.save(OUT)
print("saved v5 slides:", len(prs.slides._sldIdLst))
