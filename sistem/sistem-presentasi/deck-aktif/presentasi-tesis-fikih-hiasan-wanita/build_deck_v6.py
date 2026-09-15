# -*- coding: utf-8 -*-
"""v6 — mengikuti STRUKTUR PEMILIK verbatim: satu bagian = satu slide (tidak digabung),
pertanyaan Q1–Q3 digabung satu slide TANPA jawaban. Aspect-safe band, sanitizer, footer.
"""
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os, struct

GREEN=RGBColor(0x0F,0x3D,0x2E); GOLD=RGBColor(0xC9,0xA2,0x27)
CREAM=RGBColor(0xF7,0xF3,0xE9); INK=RGBColor(0x22,0x22,0x22)
FONT="Amiri"; W=13.333; H=7.5
HERE=os.path.dirname(os.path.abspath(__file__))
IMG=os.path.join(HERE,"gambar","pola-geometris.png")
_IW,_IH=(1376,768)
if os.path.exists(IMG):
    _d=open(IMG,'rb').read(); _IW,_IH=struct.unpack(">II",_d[16:24])
AD="٠١٢٣٤٥٦٧٨٩"
def anum(n): return "".join(AD[int(d)] for d in str(n))
def clean(t):
    return "".join(ch for ch in t if not (0x064B<=ord(ch)<=0x065F) and ord(ch)!=0x0670 and not (0x06D6<=ord(ch)<=0x06ED))

prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H)
BLANK=prs.slide_layouts[6]
def rrun(run,size,bold,color):
    run.font.name=FONT; run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=color
    rPr=run._r.get_or_add_rPr()
    for tag in ("a:ea","a:cs"):
        el=rPr.find(qn(tag))
        if el is None: el=rPr.makeelement(qn(tag),{}); rPr.append(el)
        el.set("typeface",FONT)
def rpara(p,size,bold,color,after=10,space=1.15):
    p.alignment=PP_ALIGN.RIGHT
    pPr=p._p.get_or_add_pPr(); pPr.set("rtl","1")
    p.space_after=Pt(after)
    try: p.line_spacing=space
    except Exception: pass
    if p.runs: rrun(p.runs[0],size,bold,color)
def bg(s):
    f=s.background.fill; f.solid(); f.fore_color.rgb=CREAM
def box(s,l,t,w,h): return s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
def add_band(s,bw):
    if not os.path.exists(IMG): return
    f=(bw/H)*_IH/_IW
    if f>1: f=1
    pic=s.shapes.add_picture(IMG,Inches(W-bw),0,Inches(bw),Inches(H))
    pic.crop_left=0.0; pic.crop_top=0.0; pic.crop_bottom=0.0; pic.crop_right=1-f
def new_slide(num,src,bw=0.9):
    s=prs.slides.add_slide(BLANK); bg(s); add_band(s,bw)
    gl=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(W-bw-0.06),0,Inches(0.06),Inches(H))
    gl.fill.solid(); gl.fill.fore_color.rgb=GOLD; gl.line.fill.background()
    fb=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,Inches(H-0.55),Inches(W),Inches(0.55))
    fb.fill.solid(); fb.fill.fore_color.rgb=GREEN; fb.line.fill.background()
    nb=box(s,0.5,H-0.5,2.0,0.45); p=nb.text_frame.paragraphs[0]; p.text=anum(num)
    p.alignment=PP_ALIGN.LEFT; rrun(p.runs[0],16,True,GOLD)
    sb=box(s,W-3.4,H-0.5,2.6,0.45); p=sb.text_frame.paragraphs[0]
    if src:
        p.text=src; p.alignment=PP_ALIGN.RIGHT; rrun(p.runs[0],14,False,GOLD)
    return s
def title(s,text,size=32):
    tb=box(s,0.7,0.45,W-2.2,1.4); tf=tb.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.text=clean(text); rpara(p,size,True,GREEN,after=6)
    line=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(W-0.96-6.4),Inches(1.8),Inches(6.4),Inches(0.06))
    line.fill.solid(); line.fill.fore_color.rgb=GOLD; line.line.fill.background()
def bullets(s,items,size=24):
    tb=box(s,0.9,2.0,W-2.3,4.6); tf=tb.text_frame; tf.word_wrap=True
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    first=True
    for item in items:
        p=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
        p.text=clean(item); rpara(p,size,False,INK,after=14,space=1.15)
def para(s,text,size=24):
    tb=box(s,0.9,2.1,W-2.3,4.4); tf=tb.text_frame; tf.word_wrap=True
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.text=clean(text); rpara(p,size,False,INK,after=0,space=1.4)
def notes(s,t): s.notes_slide.notes_text_frame.text=t
def table_slide(s,rows):
    sh=s.shapes.add_table(len(rows),2,Inches(1.1),Inches(2.1),Inches(W-3.4),Inches(4.3))
    t=sh.table; t.columns[0].width=Inches(3.0); t.columns[1].width=Inches(W-3.4-3.0)
    for r,(a,b) in enumerate(rows):
        for c,val in ((0,a),(1,b)):
            cell=t.cell(r,c); cell.text=clean(val)
            p=cell.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.RIGHT
            p._p.get_or_add_pPr().set("rtl","1")
            run=p.runs[0]; run.font.name=FONT; run.font.size=Pt(18 if r else 20); run.font.bold=(r==0)
            rPr=run._r.get_or_add_rPr()
            for tag in ("a:ea","a:cs"):
                el=rPr.find(qn(tag)) or rPr.makeelement(qn(tag),{})
                if el not in rPr: rPr.append(el)
                el.set("typeface",FONT)
            if r==0: cell.fill.solid(); cell.fill.fore_color.rgb=GREEN; run.font.color.rgb=CREAM
            elif r%2==0: cell.fill.solid(); cell.fill.fore_color.rgb=RGBColor(0xEC,0xE6,0xD6)
            else: cell.fill.solid(); cell.fill.fore_color.rgb=CREAM

# S1 judul
s=new_slide(1,"",bw=1.4)
tb=box(s,1.2,2.0,W-3.2,2.6); tf=tb.text_frame; tf.word_wrap=True
p=tf.paragraphs[0]; p.text=clean("أحكام النوازل الفقهية المعاصرة المتعلقة بزينة المرأة في المذاهب الأربعة")
rpara(p,38,True,GREEN,after=20,space=1.2)
p2=tf.add_paragraph(); p2.text=clean("إعداد: الشيخ علي بلعيدي"); rpara(p2,26,False,INK)
fr=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(1.0),Inches(1.7),Inches(W-3.0),Inches(3.4))
try: fr.adjustments[0]=0.04
except Exception: pass
fr.fill.background(); fr.line.color.rgb=GOLD; fr.line.width=Pt(2)
notes(s,"Pembuka. Pembimbing/universitas kosong - jangan dikarang.")

# S2 ملخص
s=new_slide(2,"ص٦"); title(s,"ملخص البحث: أحكام زينة المرأة المعاصرة في المذاهب الأربعة")
bullets(s,[
 "يتناول البحث النوازل المتعلقة بزينة المرأة: التشقير، الأظافر، الحواجب، عدسات العين، النمص، الوشم، الوصل، الباروكة، العمليات الجراحية، مساحيق الوجه وغيرها",
 "المنهج: وصفي مقارن، استقرائي، استنباطي",
 "الأصل التحريم عند تغيير خلق الله أو التغرير أو التدليس أو الضرر، أو نص الشارع",
 "التشقير إن لم يزل الشعر فليس كالنمص؛ والوشم والعدسات والأظافر تجوز للحاجة لا للتغرير",
 "الوصل ليس منه إن كان لحاجة ويحرم للتغرير؛ وتحمير الوجه يحرم على الخلية عند الشافعية",
 "العمليات تجوز للحاجة دون تغيير الخلق؛ وللزينة أثر في خيار النكاح",
],size=19)
notes(s,"ملخص hal 6 diperpanjang sesuai isi tesis.")

# S3 أهداف
s=new_slide(3,"ص٢١"); title(s,"أهداف البحث")
bullets(s,[
 "بيان حكم أدوات التجميل المعاصرة للمرأة في المذاهب الأربعة",
 "معرفة حكم عمليات التجميل في المذاهب المعتمدة",
 "معرفة حكم بيع أدوات التجميل في المذاهب الأربعة",
 "بيان الأثر المترتب على الزينة في عقد النكاح",
],size=24)
notes(s,"أهداف hal 21 (pemilik menyebut ص٢٠).")

# S4 فوائد
s=new_slide(4,"ص٢١"); title(s,"فوائد البحث")
bullets(s,[
 "معرفة أحكام الزينة المعاصرة في المذاهب الأربعة",
 "بيان أحكام النوازل المتعلقة بزينة المرأة",
 "معرفة حكم عمليات التجميل للمرأة",
 "معرفة حكم بيع أدوات التجميل الممنوعة",
 "إدراك أثر الزينة على عقد النكاح",
],size=22)
notes(s,"فوائد hal 21.")

# S5 مشكلات
s=new_slide(5,"ص٢١"); title(s,"مشكلات البحث")
bullets(s,[
 "نوازل معاصرة لم يفصلها الفقهاء قديما",
 "صعوبة حصر وسائل الزينة لاختلاط العادات",
 "اختلاف الأزمنة يستدعي نظر القواعد والضوابط",
 "ورود أدوات من الغرب دون نظر شرعي",
 "اختلاف التكييف الفقهي لوسائل الزينة",
],size=22)
notes(s,"مشكلات hal 21.")

# S6 منهج
s=new_slide(6,"ص٢٣"); title(s,"منهج البحث")
bullets(s,[
 "وصفي مقارن: تتبع آراء الفقهاء في المذاهب الأربعة من كتبهم المعتمدة، مع الرجوع للمؤلفات الحديثة وكلام المعاصرين؛ فيصف النوازل ويشخصها ويقيسها على المنصوص لاستخراج حكمها",
 "استقرائي: استقراء أحكام زينة المرأة الفقهية والطبية من كتب المذاهب القديمة والحديثة",
 "استنباطي: استنباط الأحكام المتعلقة بزينة المرأة في المذاهب الأربعة",
],size=20)
notes(s,"منهج hal 23 diperpanjang.")

# S7 أسئلة (Q1-Q3 saja, tanpa jawaban)
s=new_slide(7,"ص٩٠·١٤٦"); title(s,"أسئلة البحث")
bullets(s,[
 "السؤال الأول: ما حكم استعمال أدوات التجميل المعاصرة للمرأة في المذاهب الأربعة؟",
 "السؤال الثاني: ما حكم عمليات التجميل للمرأة في المذاهب الأربعة؟",
 "السؤال الثالث: ما حكم التعامل مع المرأة في زينتها بيعا وتجميلا وطبا في المذاهب الأربعة؟",
],size=22)
notes(s,"أسئلة dari Bab 4, HANYA soal tanpa jawaban, digabung satu slide. Q1 verbatim ص٩٠; Q3 verbatim ص١٤٦; Q2 direkonstruksi dari tujuan (heading Bab 4 ~ص١٤٢ belum terbaca verbatim) - tunggu tinjauan pemilik.")

# S8 العلة
s=new_slide(8,"ص٦٦–٦٧"); title(s,"اتفقوا على الجواز بغير شعر الآدمي واختلفوا في العلة، والراجح التدليس")
table_slide(s,[
 ("المذهب","العلة في تحريم وصل الشعر"),
 ("الحنفية","كرامة الإنسان وتحريم الابتذال"),
 ("المالكية","مركبة: تغيير خلق الله والتدليس"),
 ("الشافعية","مركبة: الكرامة وتحريم استعمال النجس"),
 ("الحنابلة","التدليس والغش"),
 ("الراجح","الزور والتدليس؛ لحديث المرأة"),
])
notes(s,"العلة hal 66-67.")

# S9 جنس ونوع
s=new_slide(9,"ص٨٨"); title(s,"جنس البحث ونوعه")
bullets(s,[
 "جنس البحث: فقهي مقارن؛ يتتبع آراء الفقهاء ويشخص النوازل ويقيسها على المسائل المنصوص عليها",
 "نوعه: استقرائي لاستخراج قواعد عامة تساعد على فهم النوازل؛ ثم استنباطي لتنزيل الأحكام على المسائل المعاصرة",
 "اختيرت المناهج لمناسبتها الموضوع، وللاطلاع على كتب الفقه والحديث والتفسير ومعاجم اللغة والدراسات المعاصرة",
],size=20)
notes(s,"جنس ونوع hal 88 diperpanjang.")

# S10 جمع
s=new_slide(10,"ص٨٩"); title(s,"طريقة جمع المعلومات")
bullets(s,[
 "تتبع آراء الفقهاء في المذاهب الأربعة واستقراؤها والاطلاع على المسائل المتعلقة بالزينة",
 "تتبع محال ذكرها في أبواب كثيرة: اللباس، النفقات، العيدين، الجمعة، الجنائز، ستر العورة، ووصل الشعر وغيره",
 "الاطلاع على البحوث والدراسات وكتب الفتاوى، ثم ترتيب المعلومات وإلحاق المسائل المعاصرة بأحكامها، مرجحا قولا على قول",
],size=20)
notes(s,"جمع hal 89 diperpanjang.")

# S11 معالجة
s=new_slide(11,"ص٨٩"); title(s,"معالجة المعلومات")
bullets(s,[
 "بعد الاطلاع على كتب الفقه والفتاوى والتفسير والحديث وشروحه ومعاجم اللغة، والمقارنة بين آراء المذاهب",
 "ذكر الأدلة ومناقشتها مع ترجيح الراجح، وإلحاق المسائل المعاصرة بأصولها المنصوص عليها",
 "الاعتماد على نصوص الفقهاء في كتبهم المعتمدة: فقهية أو حديثية أو تفسيرية أو فتاوى",
],size=20)
notes(s,"معالجة hal 89 diperpanjang.")

# S12 تحقق
s=new_slide(12,"ص٨٩"); title(s,"التحقق من صحة المعلومات")
bullets(s,[
 "تحقق الباحث من المعلومات المأخوذة بعرضها على عدة كتب ومصادر مختلفة",
 "اعتمادها عند التوصل إلى الاتفاق في المعنى وإن اختلفت الألفاظ",
])
notes(s,"تحقق hal 89 diperpanjang.")

# S13 خاتمة
s=new_slide(13,"ص١٥١"); title(s,"الخاتمة")
bullets(s,[
 "الحمد لله على إتمام البحث",
 "جهد بشري يقبل الخطأ والصواب",
 "نسأل الله القبول والنفع به",
])
notes(s,"خاتمة hal 151 (pemilik menyebut ص١٥٢).")

# S14 توصيات
s=new_slide(14,"ص١٥٢"); title(s,"التوصيات")
bullets(s,[
 "للعلماء وطلاب العلم: دقة البحث وإظهار أحكام النوازل للمجتمع",
 "للمرأة المسلمة: الحذر من الزينة المحرمة والاكتفاء بالحلال عن علم",
 "للمجامع الفقهية: متابعة المستجدات في زينة المرأة وبيان أحكامها",
])
notes(s,"توصيات hal 152 (pemilik menyebut ص١٥٣).")

OUT=os.path.join(HERE,"keluaran","presentasi-tesis-fikih-hiasan-wanita.pptx")
prs.save(OUT)
print("saved v6 slides:", len(prs.slides._sldIdLst))
